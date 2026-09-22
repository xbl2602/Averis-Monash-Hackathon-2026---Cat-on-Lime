#!/usr/bin/env python3
"""Office-file generator for the perturbation test set (called by scripts/perturb-generate.mjs).

Input: a JSON task file, each item {kind: "xlsx"|"docx", outPath, title, rows:[{label,value}], notes:[str]}
Purpose: rebuild txt attachments as "weird but valid" xlsx / docx files, to test the parser's
robustness against layout variation.
"""
import json
import sys
from pathlib import Path


def write_xlsx(task):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet("DATA")
    else:
        ws.title = "DATA"
    ws.merge_cells("A1:D2")
    ws["A1"] = task.get("title", "")
    row = 4
    for index, item in enumerate(task["rows"]):
        ws.cell(row=row, column=1, value=item["label"])
        ws.cell(row=row, column=2, value=item["value"])
        row += 2 if index % 2 == 0 else 1
    notes = wb.create_sheet("NOTES")
    for line_no, note in enumerate(task.get("notes", []), start=1):
        notes.cell(row=line_no, column=1, value=note)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 60
    wb.save(task["outPath"])


def write_docx(task):
    from docx import Document

    doc = Document()
    doc.add_heading(task.get("title", "DOCUMENT"), level=1)
    doc.add_paragraph("")
    table = doc.add_table(rows=0, cols=2)
    for item in task["rows"]:
        cells = table.add_row().cells
        cells[0].text = item["label"]
        cells[1].text = item["value"]
    doc.add_paragraph("")
    for note in task.get("notes", []):
        doc.add_paragraph(note)
    doc.save(task["outPath"])


def main():
    tasks = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    for task in tasks:
        Path(task["outPath"]).parent.mkdir(parents=True, exist_ok=True)
        if task["kind"] == "xlsx":
            write_xlsx(task)
        elif task["kind"] == "docx":
            write_docx(task)
        else:
            raise SystemExit(f"unknown kind: {task['kind']}")
    print(f"office files written: {len(tasks)}")


if __name__ == "__main__":
    main()
